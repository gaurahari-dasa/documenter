import re
from docx import Document


def replace_text_in_docx(doc: Document, placeholders: dict):
    for para in doc.paragraphs:
        for key, value in placeholders.items():
            if not value:
                # Construct regex pattern dynamically based on the tag
                pattern = re.compile(fr"\{{{re.escape(key)}: .*? :{re.escape(key)}\}}")
                para.text = re.sub(pattern, '', para.text)
            else:
                pattern = re.compile(fr"\{{{re.escape(key)}: (.*?) :{re.escape(key)}\}}")
                para.text = re.sub(pattern, r'\1', para.text)

    # doc.save(new_file_path)
